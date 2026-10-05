"""
prism_data_pull.py
--------------------
Workflow: you maintain an Excel roster with a "Session ID" column filled in
for each row (and a "Session Group ID" column for session-group-session
rows — Prism's API doesn't return that ID, same as it doesn't for the
group_session.py mutation, so it has to come from you). Run this script and
it looks up each Session ID against Prism and writes the rest of the
columns back into the SAME file:

    Session Name, Starts At, Ends At, Status, Join URL, Start URL,
    Session Type, Meeting Code

"Session Type" (session / workshop-session / session-group-session) is
looked up automatically from Prism's `virtualMeetingFor` field — you don't
need to know it upfront, this fills it in for you. It's the value
Zoom_Ai_V54.py's "Prism Type" column and prism_status_push.py both expect.

USAGE
-----
    python prism_data_pull.py roster.xlsx
    python prism_data_pull.py roster.xlsx --sheet "Sheet1"

Requires PRISM_BEARER_TOKEN — set as an env var, or pasted into
prism_config.py (see that file for why it's git-ignored).

WHAT IT DOES NOT TOUCH
-----------------------
- "Session ID" / "Session Group ID" — these are your input, never
  overwritten.
- Any other columns already in the sheet (Program Name, Moderator, Host
  Email, etc.) — untouched, so you can run this on the same roster you'll
  eventually feed into Command Center as long as the header names match.
"""

import os
import argparse
import asyncio
import requests
from openpyxl import load_workbook

from prism_config import PRISM_BEARER_TOKEN

PRISM_GRAPHQL_URL = "https://lxp-api.upgrad.com/graphql"

_QUERY = """
query liveSessions($where: LiveSessionWhereInput, $sort: LiveSessionSortInput, $limit: Int, $skip: Int) {
  liveSessions(where: $where, skip: $skip, limit: $limit, sort: $sort) {
    totalCount
    result {
      id
      session {
        id
        name
        status
        startsAt
        endsAt
        virtualMeeting { code joinUrl startUrl virtualMeetingFor }
      }
      workshopSession {
        id
        name
        status
        startsAt
        endsAt
        virtualMeeting { code joinUrl startUrl virtualMeetingFor }
      }
    }
  }
}
"""

OUTPUT_COLUMNS = ["Session Name", "Starts At", "Ends At", "Status",
                   "Join URL", "Start URL", "Session Type", "Meeting Code"]


def _headers():
    if not PRISM_BEARER_TOKEN:
        raise RuntimeError(
            "No Prism bearer token found. Set PRISM_BEARER_TOKEN as an env "
            "var, or paste it into prism_config.py."
        )
    return {
        "accept": "*/*",
        "apollographql-client-name": "prism-frontend",
        "apollographql-client-version": "1.0",
        "authorization": f"Bearer {PRISM_BEARER_TOKEN}",
        "content-type": "application/json",
        "role": "delivery-manager",
    }


def fetch_sessions(session_ids):
    """Returns {session_id: {output columns...}} for every ID Prism recognized."""
    if not session_ids:
        return {}
    payload = {
        "query": _QUERY,
        "variables": {
            "limit": len(session_ids),
            "skip": 0,
            "where": {"sessions": {"_in": session_ids}},
            "sort": {"startsAt": "desc"},
        },
    }
    resp = requests.post(PRISM_GRAPHQL_URL, headers=_headers(), json=payload, timeout=30)
    try:
        resp.raise_for_status()
    except requests.exceptions.HTTPError as e:
        # requests' default HTTPError only carries the status line ("400
        # Client Error: Bad Request for url: ...") — the actual reason
        # (expired/invalid token, bad signature, missing role header,
        # etc.) is in the response BODY, which raise_for_status() throws
        # away. Re-raise with that body attached so the Event Log / file
        # log actually tells you why, instead of just that it happened.
        raise RuntimeError(
            f"Prism API HTTP {resp.status_code}: {resp.text[:500]}"
        ) from e
    data = resp.json()
    if "errors" in data:
        raise RuntimeError(f"Prism API error: {data['errors']}")

    by_id = {}
    for item in data["data"]["liveSessions"]["result"]:
        s  = item.get("session") or {}
        ws = item.get("workshopSession") or {}
        entity = s if s else ws          # mutually exclusive per item
        vm = entity.get("virtualMeeting") or {}
        entity_id = entity.get("id")
        if not entity_id:
            continue
        by_id[entity_id] = {
            "Session Name": entity.get("name", ""),
            "Starts At":    entity.get("startsAt", ""),
            "Ends At":      entity.get("endsAt", ""),
            "Status":       entity.get("status", ""),
            "Join URL":     vm.get("joinUrl", ""),
            "Start URL":    vm.get("startUrl", ""),
            "Session Type": vm.get("virtualMeetingFor", ""),
            "Meeting Code": vm.get("code", ""),
        }
    return by_id


async def fetch_sessions_async(session_ids):
    """Async-safe wrapper — same reason prism_status_push.py has one:
    requests.post() blocks, and the app's engine + UI share one event
    loop (qasync), so a blocking call here would freeze the whole window
    for the duration of the request."""
    return await asyncio.to_thread(fetch_sessions, session_ids)


def _find_col(header_row, *names):
    """Case-insensitive header lookup. Returns 1-based column index or None."""
    lowered = {str(h).strip().lower(): i + 1 for i, h in enumerate(header_row) if h}
    for name in names:
        if name.lower() in lowered:
            return lowered[name.lower()]
    return None


def update_excel_in_place(path, sheet_name=None):
    wb = load_workbook(path)
    ws = wb[sheet_name] if sheet_name else wb.active

    header_row = [cell.value for cell in ws[1]]
    id_col = _find_col(header_row, "Session ID", "Prism Session ID")
    if not id_col:
        raise RuntimeError(
            "No 'Session ID' column found in the header row (row 1). "
            "Add that column and fill in the Prism session IDs first."
        )

    # Ensure every output column exists — append any that are missing,
    # right after the current last column, without disturbing existing ones.
    out_cols = {}
    next_free_col = len(header_row) + 1
    for name in OUTPUT_COLUMNS:
        col = _find_col(header_row, name)
        if not col:
            col = next_free_col
            ws.cell(row=1, column=col, value=name)
            header_row.append(name)
            next_free_col += 1
        out_cols[name] = col

    # Collect Session IDs from every data row (a Session ID can repeat if
    # you keep multiple rows per session — all matching rows get updated).
    row_for_id = {}
    for r in range(2, ws.max_row + 1):
        sid = ws.cell(row=r, column=id_col).value
        if sid:
            row_for_id.setdefault(str(sid).strip(), []).append(r)

    unique_ids = sorted(row_for_id.keys())
    if not unique_ids:
        print("No Session IDs found in the sheet — nothing to do.")
        return

    print(f"Looking up {len(unique_ids)} session ID(s) against Prism...")
    by_id = fetch_sessions(unique_ids)

    found = missing = 0
    for sid, rows in row_for_id.items():
        data = by_id.get(sid)
        for r in rows:
            if data:
                for name, col in out_cols.items():
                    ws.cell(row=r, column=col, value=data[name])
                found += 1
            else:
                ws.cell(row=r, column=out_cols["Status"], value="NOT FOUND IN PRISM")
                missing += 1

    wb.save(path)
    print(f"Updated {path} — {found} row(s) matched, {missing} row(s) not found in Prism.")
    if missing:
        print("Rows marked 'NOT FOUND IN PRISM' — double check those Session IDs.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Fill in Prism session data (name/times/status/URLs/type) "
                     "for every row with a Session ID, in place.")
    parser.add_argument("excel_path", help="Path to the .xlsx roster to update.")
    parser.add_argument("--sheet", default=None, help="Sheet name (default: active sheet).")
    args = parser.parse_args()
    try:
        update_excel_in_place(args.excel_path, args.sheet)
    except PermissionError:
        print(f"Could not save {args.excel_path} — it's probably still open in Excel. "
              f"Close it and run this again.")
    except requests.exceptions.RequestException as e:
        print(f"Could not reach Prism ({PRISM_GRAPHQL_URL}): {e}")
    except RuntimeError as e:
        print(f"Error: {e}")
